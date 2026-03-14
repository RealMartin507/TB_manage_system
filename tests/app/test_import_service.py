"""
导入逻辑单元测试 - Phase 3
覆盖：
1. 新图斑导入 → 正确新增 tb_plot + tb_plot_task
2. 已有图斑重复导入 → 只追加 tb_plot_task，不覆盖图斑基本信息
3. 空值字段 → 存为 None，不报错
4. mapping 中有「忽略」的列 → 跳过，不写入数据库
5. report_status 未映射 → 默认存为「待报」
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

import openpyxl
import pytest

from app.services.import_service import (
    IGNORE_VALUE,
    execute_import,
    import_plot_task_batch,
    parse_excel,
    save_upload,
)

# ──────────────────────────────────────────────
# Fixtures
# ──────────────────────────────────────────────

_SCHEMA = """
CREATE TABLE IF NOT EXISTS tb_plot (
    plot_id       TEXT    PRIMARY KEY,
    township      TEXT,
    village       TEXT,
    area          REAL,
    basic_farmland REAL,
    plot_type     TEXT,
    land_before   TEXT,
    land_after    TEXT,
    folder_path   TEXT,
    created_at    DATETIME DEFAULT (datetime('now', 'localtime'))
);

CREATE TABLE IF NOT EXISTS tb_platform (
    platform_id   INTEGER PRIMARY KEY AUTOINCREMENT,
    platform_name TEXT    NOT NULL,
    platform_url  TEXT,
    sort_order    INTEGER DEFAULT 0
);

CREATE TABLE IF NOT EXISTS tb_task (
    task_id       INTEGER PRIMARY KEY AUTOINCREMENT,
    task_name     TEXT    NOT NULL,
    source        TEXT,
    platform_id   INTEGER REFERENCES tb_platform(platform_id) ON DELETE SET NULL,
    created_time  DATE,
    deadline      DATE,
    import_time   DATETIME DEFAULT (datetime('now', 'localtime'))
);

CREATE TABLE IF NOT EXISTS tb_plot_task (
    id                INTEGER PRIMARY KEY AUTOINCREMENT,
    plot_id           TEXT    NOT NULL REFERENCES tb_plot(plot_id) ON DELETE CASCADE,
    task_id           INTEGER NOT NULL REFERENCES tb_task(task_id) ON DELETE CASCADE,
    seq_no            INTEGER,
    project_name      TEXT,
    land_unit         TEXT,
    approval_doc      TEXT,
    land_supply_doc   TEXT,
    legality          TEXT,
    internal_review   TEXT,
    field_survey      TEXT,
    report_status     TEXT DEFAULT '待报',
    reject_reason     TEXT,
    archive_status    TEXT,
    dispatch          TEXT,
    overlap_situation TEXT,
    overlap_ratio     REAL,
    overlap_area      REAL,
    remark            TEXT
);
"""


@pytest.fixture
def db():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(_SCHEMA)
    conn.commit()
    yield conn
    conn.close()


def _make_session(tmp_path: Path, rows: list[list]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    for row in rows:
        ws.append(row)
    file_path = tmp_path / "test.xlsx"
    wb.save(file_path)
    result = save_upload(file_path.read_bytes(), "test.xlsx")
    return result["session_id"]


def _do_import(db: sqlite3.Connection, session_id: str, mapping: dict[str, str],
               source: str = "来源", task_name: str = "任务",
               data_start: int = 2, data_end: int = 2):
    return execute_import(
        db,
        session_id=session_id,
        sheet_name="Sheet",
        header_row=1,
        data_start_row=data_start,
        data_end_row=data_end,
        source=source,
        task_name=task_name,
        platform_id=None,
        mapping=mapping,
    )


# ──────────────────────────────────────────────
# 1. 新图斑导入
# ──────────────────────────────────────────────

class TestNewPlotImport:
    def test_inserts_plot_and_plot_task(self, db: sqlite3.Connection, tmp_path: Path):
        """新图斑导入后，tb_plot 和 tb_plot_task 各增加一条记录。"""
        sid = _make_session(tmp_path, [
            ["监测编号", "乡镇名称", "村名称", "填报"],
            ["P001",    "东城镇",   "新村",  "待报"],
        ])

        result = _do_import(db, sid, {
            "plot_id":       "监测编号",
            "township":      "乡镇名称",
            "village":       "村名称",
            "report_status": "填报",
        })

        assert result["added"] == 1
        assert result["skipped"] == 0

        plot = db.execute("SELECT * FROM tb_plot WHERE plot_id = 'P001'").fetchone()
        assert plot is not None
        assert plot["township"] == "东城镇"
        assert plot["village"] == "新村"

        pt = db.execute("SELECT * FROM tb_plot_task WHERE plot_id = 'P001'").fetchone()
        assert pt is not None
        assert pt["report_status"] == "待报"

    def test_creates_task_if_not_exists(self, db: sqlite3.Connection, tmp_path: Path):
        """导入时自动创建不存在的任务记录。"""
        sid = _make_session(tmp_path, [["监测编号"], ["P002"]])

        _do_import(db, sid, {"plot_id": "监测编号"}, source="来源A", task_name="任务A")

        task = db.execute("SELECT * FROM tb_task WHERE task_name = '任务A'").fetchone()
        assert task is not None
        assert task["source"] == "来源A"


# ──────────────────────────────────────────────
# 2. 重复导入已有图斑
# ──────────────────────────────────────────────

class TestDuplicatePlotImport:
    def test_does_not_overwrite_plot(self, db: sqlite3.Connection, tmp_path: Path):
        """已有图斑重复导入时，tb_plot 数据不被覆盖，只追加 tb_plot_task。"""
        db.execute(
            "INSERT INTO tb_plot (plot_id, township, village) VALUES ('P001', '原乡镇', '原村')"
        )
        db.commit()

        sid = _make_session(tmp_path, [
            ["监测编号", "乡镇名称"],
            ["P001",    "新乡镇"],
        ])

        result = _do_import(db, sid, {"plot_id": "监测编号", "township": "乡镇名称"})

        assert result["added"] == 0
        assert result["updated"] == 1

        plot = db.execute("SELECT * FROM tb_plot WHERE plot_id = 'P001'").fetchone()
        assert plot["township"] == "原乡镇"

    def test_appends_plot_task_for_different_tasks(self, db: sqlite3.Connection, tmp_path: Path):
        """同一图斑不同任务各自生成 plot_task 记录。"""
        db.execute("INSERT INTO tb_plot (plot_id) VALUES ('P001')")
        db.commit()

        for i in range(2):
            sid = _make_session(tmp_path, [["监测编号"], ["P001"]])
            _do_import(db, sid, {"plot_id": "监测编号"}, source=f"来源{i}", task_name=f"任务{i}")

        pt_count = db.execute(
            "SELECT COUNT(*) FROM tb_plot_task WHERE plot_id = 'P001'"
        ).fetchone()[0]
        assert pt_count == 2

    def test_no_duplicate_plot_task_same_task(self, db: sqlite3.Connection, tmp_path: Path):
        """同一图斑同一任务重复导入，不重复写 plot_task。"""
        db.execute("INSERT INTO tb_plot (plot_id) VALUES ('P001')")
        db.commit()

        for _ in range(2):
            sid = _make_session(tmp_path, [["监测编号"], ["P001"]])
            _do_import(db, sid, {"plot_id": "监测编号"}, source="来源", task_name="任务")

        pt_count = db.execute(
            "SELECT COUNT(*) FROM tb_plot_task WHERE plot_id = 'P001'"
        ).fetchone()[0]
        assert pt_count == 1


# ──────────────────────────────────────────────
# 3. 空值字段 → None，不报错
# ──────────────────────────────────────────────

class TestNullFields:
    def test_empty_string_cells_stored_as_none(self, db: sqlite3.Connection, tmp_path: Path):
        """Excel 空字符串单元格应存为 NULL，不报错。"""
        sid = _make_session(tmp_path, [
            ["监测编号", "乡镇名称", "监测面积"],
            ["P001",    "",        ""],
        ])

        _do_import(db, sid, {
            "plot_id":  "监测编号",
            "township": "乡镇名称",
            "area":     "监测面积",
        })

        plot = db.execute("SELECT * FROM tb_plot WHERE plot_id = 'P001'").fetchone()
        assert plot["township"] is None
        assert plot["area"] is None

    def test_none_cell_stored_as_none(self, db: sqlite3.Connection, tmp_path: Path):
        """openpyxl 返回 None 的单元格也应正确处理。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws.append(["监测编号", "乡镇名称"])
        ws.append(["P002", None])
        file_path = tmp_path / "null_test.xlsx"
        wb.save(file_path)

        result = save_upload(file_path.read_bytes(), "null_test.xlsx")
        _do_import(db, result["session_id"], {"plot_id": "监测编号", "township": "乡镇名称"})

        plot = db.execute("SELECT * FROM tb_plot WHERE plot_id = 'P002'").fetchone()
        assert plot["township"] is None


# ──────────────────────────────────────────────
# 4. mapping 中「忽略」列 → 不写入数据库
# ──────────────────────────────────────────────

class TestIgnoreMapping:
    def test_ignored_column_not_stored(self, db: sqlite3.Connection, tmp_path: Path):
        """mapping 值为 __ignore__ 的字段不应写入数据库。"""
        sid = _make_session(tmp_path, [
            ["监测编号", "乡镇名称", "备注"],
            ["P001",    "东城镇",   "这是备注"],
        ])

        _do_import(db, sid, {
            "plot_id":  "监测编号",
            "township": "乡镇名称",
            "remark":   IGNORE_VALUE,
        })

        pt = db.execute("SELECT * FROM tb_plot_task WHERE plot_id = 'P001'").fetchone()
        assert pt["remark"] is None


# ──────────────────────────────────────────────
# 5. report_status 未映射 → 默认「待报」
# ──────────────────────────────────────────────

class TestReportStatusDefault:
    def test_default_when_not_in_mapping(self, db: sqlite3.Connection, tmp_path: Path):
        """mapping 中没有 report_status 时，数据库应存为「待报」。"""
        sid = _make_session(tmp_path, [["监测编号"], ["P001"]])
        _do_import(db, sid, {"plot_id": "监测编号"})

        pt = db.execute("SELECT * FROM tb_plot_task WHERE plot_id = 'P001'").fetchone()
        assert pt["report_status"] == "待报"

    def test_default_when_ignored(self, db: sqlite3.Connection, tmp_path: Path):
        """report_status 映射为忽略时，数据库应存为「待报」。"""
        sid = _make_session(tmp_path, [
            ["监测编号", "填报"],
            ["P001",    "已报"],
        ])

        _do_import(db, sid, {
            "plot_id":       "监测编号",
            "report_status": IGNORE_VALUE,
        })

        pt = db.execute("SELECT * FROM tb_plot_task WHERE plot_id = 'P001'").fetchone()
        assert pt["report_status"] == "待报"


# ──────────────────────────────────────────────
# 兼容旧接口测试
# ──────────────────────────────────────────────

class TestLegacyInterface:
    def test_parse_excel_basic(self, tmp_path: Path):
        """parse_excel 正常解析并按 mapping 提取字段。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["监测编号", "乡镇名称"])
        ws.append(["P001", "东城镇"])
        file_path = tmp_path / "legacy.xlsx"
        wb.save(file_path)

        records = parse_excel(str(file_path), {"plot_id": "监测编号", "township": "乡镇名称"})
        assert len(records) == 1
        assert records[0]["plot_id"] == "P001"
        assert records[0]["township"] == "东城镇"

    def test_parse_excel_empty(self, tmp_path: Path):
        """parse_excel 空文件返回空列表。"""
        wb = openpyxl.Workbook()
        file_path = tmp_path / "empty.xlsx"
        wb.save(file_path)
        records = parse_excel(str(file_path), {"plot_id": "监测编号"})
        assert records == []

    def test_import_plot_task_batch(self, db: sqlite3.Connection):
        """import_plot_task_batch 批量写入。"""
        db.execute("INSERT INTO tb_plot (plot_id) VALUES ('P001')")
        db.execute("INSERT INTO tb_task (task_name, source) VALUES ('任务', '来源')")
        task_id = db.execute("SELECT task_id FROM tb_task").fetchone()["task_id"]
        db.commit()

        result = import_plot_task_batch(
            db,
            [{"plot_id": "P001", "report_status": "待报"}],
            task_id=task_id,
        )
        assert result["added"] == 0
        assert result["updated"] == 1

        pt = db.execute("SELECT * FROM tb_plot_task WHERE plot_id = 'P001'").fetchone()
        assert pt is not None
        assert pt["report_status"] == "待报"
